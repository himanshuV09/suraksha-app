import io
import zipfile
import base64
import mimetypes

from docx import Document
from pptx import Presentation


class OfficePreview:

    # --------------------------------------------------------
    # DOCX
    # --------------------------------------------------------

    def preview_docx(self, evidence, data: bytes):

        document = Document(io.BytesIO(data))

        text = "\n".join(
            p.text
            for p in document.paragraphs
        )

        tables = []

        for table in document.tables:

            rows = []

            for row in table.rows:

                rows.append([
                    cell.text
                    for cell in row.cells
                ])

            tables.append(rows)

        images = self._extract_images(
            data,
            "word/media/"
        )

        props = document.core_properties

        metadata = {
            "author": props.author,
            "title": props.title,
            "subject": props.subject,
            "created": str(props.created),
            "modified": str(props.modified),
            "last_modified_by": props.last_modified_by,
        }

        return {

            "success": True,

            "filename": evidence.filename,

            "size": evidence.file_size,

            "mime": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",

            "type": "docx",

            "metadata": metadata,

            "text": text,

            "tables": tables,

            "images": images

        }

    # --------------------------------------------------------
    # PPTX
    # --------------------------------------------------------

    def preview_pptx(self, evidence,  data: bytes):

        presentation = Presentation(
            io.BytesIO(data)
        )

        slides = []

        for slide_number, slide in enumerate(
            presentation.slides,
            start=1
        ):

            texts = []

            for shape in slide.shapes:

                if hasattr(shape, "text"):

                    if shape.text.strip():

                        texts.append(
                            shape.text
                        )

            slides.append({

                "slide": slide_number,

                "text": "\n".join(texts)

            })

        images = self._extract_images(
            data,
            "ppt/media/"
        )

        props = presentation.core_properties

        metadata = {

            "author": props.author,

            "title": props.title,

            "subject": props.subject,

            "created": str(props.created),

            "modified": str(props.modified),

            "last_modified_by": props.last_modified_by,

        }

        return {

            "success": True,

            "filename": evidence.filename,

            "size": evidence.file_size,

            "mime": "application/vnd.openxmlformats-officedocument.presentationml.presentation",

            "type": "pptx",

            "metadata": metadata,

            "slides": slides,

            "images": images

        }

    # --------------------------------------------------------
    # Common Image Extractor
    # --------------------------------------------------------

    def _extract_images(
        self,
        data: bytes,
        media_folder: str
    ):

        archive = zipfile.ZipFile(
            io.BytesIO(data)
        )

        images = []

        for filename in archive.namelist():

            if not filename.startswith(
                media_folder
            ):
                continue

            image = archive.read(filename)

            mime, _ = mimetypes.guess_type(
                filename
            )

            if mime is None:

                mime = "application/octet-stream"

            images.append({

                "name": filename.split("/")[-1],

                "mime": mime,

                "size": len(image),

                "previewUrl":

                    f"data:{mime};base64,"

                    + base64.b64encode(
                        image
                    ).decode("ascii")

            })

        return images